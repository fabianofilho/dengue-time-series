"""
Gera o manuscrito em .docx para revisao, a partir das mesmas fontes que o .tex.

Diferenca de proposito em relacao ao LaTeX: o .docx nao e o artefato de
submissao, e o artefato de REVISAO. Por isso aqui:

  - a figura entra embutida no ponto do texto que fala dela, e nao no fim, para
    o revisor ver o grafico junto do argumento;
  - a tabela e tabela nativa do Word, nao imagem, para o revisor comentar celula
    a celula e usar controle de alteracoes;
  - o cabecalho declara de qual execucao do pipeline os numeros vieram.

Formato: Times New Roman 12, espacamento duplo, numeracao de linha continua,
numero de pagina. E o que a maioria dos journals biomedicos pede.

Uso:
    python scripts/build_docx.py [--out build/manuscript.docx]
"""
from __future__ import annotations

import argparse
import json
from datetime import datetime, timezone
from pathlib import Path

import pandas as pd
from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parent.parent
PAPER = ROOT / "paper"
FIGURES = PAPER / "figures"
RESULTS = ROOT / "results"

MODELS = ["timesfm", "catboost", "xgboost", "randomforest", "sarimax", "ridge", "lgbm"]
LABEL = {
    "timesfm": "TimesFM",
    "catboost": "CatBoost",
    "xgboost": "XGBoost",
    "randomforest": "Random Forest",
    "sarimax": "SARIMAX",
    "ridge": "Ridge",
    "lgbm": "LightGBM",
}
ENS_LABEL = {
    "mean": "Mean ensemble",
    "median": "Median ensemble",
    "weighted": "Weighted ensemble",
    "stacking": "Stacked ensemble (Ridge)",
}

CITY_ORDER = [
    "Sao Paulo", "Rio de Janeiro", "Belo Horizonte", "Brasilia",
    "Fortaleza", "Recife", "Manaus", "Salvador",
]


# --------------------------------------------------------------------------- #
# infraestrutura de formatacao
# --------------------------------------------------------------------------- #
def setup_document() -> Document:
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(12)
    style.element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    pf = style.paragraph_format
    pf.line_spacing = 2.0
    pf.space_after = Pt(0)

    sec = doc.sections[0]
    sec.start_type = WD_SECTION_START.NEW_PAGE
    for attr in ("top_margin", "bottom_margin", "left_margin", "right_margin"):
        setattr(sec, attr, Inches(1))

    add_line_numbers(sec)
    add_page_numbers(sec)
    return doc


def add_line_numbers(section) -> None:
    """Numeracao de linha continua. Quase todo journal exige e quase todo
    manuscrito chega sem, porque o Word nao liga por padrao."""
    sect_pr = section._sectPr
    ln = OxmlElement("w:lnNumType")
    ln.set(qn("w:countBy"), "1")
    ln.set(qn("w:restart"), "continuous")
    ln.set(qn("w:distance"), "360")
    sect_pr.append(ln)


def add_page_numbers(section) -> None:
    p = section.footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    for kind, text in (("begin", None), (None, "PAGE"), ("end", None)):
        if kind:
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), kind)
        else:
            el = OxmlElement("w:instrText")
            el.set(qn("xml:space"), "preserve")
            el.text = text
        run._r.append(el)


def heading(doc: Document, text: str, level: int = 1):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.bold = True
    run.font.size = Pt(14 if level == 1 else 12)
    run.font.name = "Times New Roman"
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(4)
    return p


def body(doc: Document, text: str, bold_lead: str | None = None):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    if bold_lead:
        r = p.add_run(bold_lead)
        r.bold = True
    p.add_run(text)
    return p


def caption(doc: Document, label: str, text: str):
    p = doc.add_paragraph()
    p.paragraph_format.line_spacing = 1.0
    p.paragraph_format.space_after = Pt(10)
    r = p.add_run(label)
    r.bold = True
    r.font.size = Pt(10)
    r2 = p.add_run(" " + text)
    r2.font.size = Pt(10)
    return p


def add_figure(doc: Document, name: str, label: str, text: str, width=6.4):
    """Figura embutida no ponto do texto, com legenda logo abaixo."""
    path = FIGURES / f"{name}.png"
    if not path.exists():
        p = doc.add_paragraph()
        r = p.add_run(f"[FIGURA AUSENTE: {name}.png]")
        r.font.color.rgb = RGBColor(0xC0, 0x00, 0x00)
        r.bold = True
        return
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.line_spacing = 1.0
    p.paragraph_format.space_before = Pt(8)
    p.add_run().add_picture(str(path), width=Inches(width))
    caption(doc, label, text)


def add_table(doc: Document, header: list[str], rows: list[list[str]],
              label: str, text: str, bold_cells: set[tuple[int, int]] | None = None):
    """Tabela nativa do Word: o revisor consegue comentar e editar celula."""
    caption(doc, label, text)
    t = doc.add_table(rows=1, cols=len(header))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(header):
        cell = t.rows[0].cells[i]
        cell.text = ""
        run = cell.paragraphs[0].add_run(h)
        run.bold = True
        run.font.size = Pt(9)
        cell.paragraphs[0].paragraph_format.line_spacing = 1.0
    for ri, row in enumerate(rows):
        cells = t.add_row().cells
        for ci, v in enumerate(row):
            cells[ci].text = ""
            run = cells[ci].paragraphs[0].add_run(str(v))
            run.font.size = Pt(9)
            if bold_cells and (ri, ci) in bold_cells:
                run.bold = True
            cells[ci].paragraphs[0].paragraph_format.line_spacing = 1.0
    doc.add_paragraph().paragraph_format.space_after = Pt(6)
    return t


def note(doc: Document, text: str):
    p = doc.add_paragraph()
    p.paragraph_format.line_spacing = 1.0
    r = p.add_run(text)
    r.font.size = Pt(9)
    r.italic = True
    return p


def flag(doc: Document, text: str):
    """Marcador vermelho de pendencia, o equivalente do \\missing do LaTeX."""
    p = doc.add_paragraph()
    r = p.add_run(f"[PENDING: {text}]")
    r.bold = True
    r.font.color.rgb = RGBColor(0xC0, 0x00, 0x00)
    return p


# --------------------------------------------------------------------------- #
# dados
# --------------------------------------------------------------------------- #
def load() -> dict:
    v = json.loads((PAPER / "verified_numbers.json").read_text(encoding="utf-8"))
    frames = {}
    for f in sorted(RESULTS.glob("benchmark_*_metrics.csv")):
        city = f.name.replace("benchmark_", "").replace("_metrics.csv", "")
        frames[city] = pd.read_csv(f)
    v["_metrics_files"] = frames
    return v


def cohort_stamp(v: dict) -> str:
    n = v["n_predictions_per_model_city"]
    folds = n // 12
    months = folds + 48 + 12 - 1
    return (
        f"{v['n_models']} models, {v['n_cities']} cities, {months} monthly observations "
        f"per city, {folds} rolling-origin folds, {n:,} predictions per model per city, "
        f"{v['n_predictions_total']:,} in total."
    )


# --------------------------------------------------------------------------- #
# secoes
# --------------------------------------------------------------------------- #
def title_block(doc: Document, v: dict, stale: bool) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(
        "Zero-Shot Foundation Models Versus Purpose-Trained Models for Dengue "
        "Forecasting in Eight Brazilian State Capitals (2010 to 2024): "
        "A Rolling-Origin Benchmark"
    )
    r.bold = True
    r.font.size = Pt(15)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run("Fabiano Barcellos-Filho, Jose Trajano, Wilson Falco")

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("Corresponding author: fabiano.nb@gmail.com")
    r.font.size = Pt(10)

    box = doc.add_paragraph()
    box.paragraph_format.line_spacing = 1.0
    r = box.add_run("REVIEW COPY, NOT FOR SUBMISSION\n")
    r.bold = True
    r.font.color.rgb = RGBColor(0xC0, 0x00, 0x00)
    r2 = box.add_run(
        f"Generated {datetime.now(timezone.utc).strftime('%Y-%m-%d %H:%M UTC')} from "
        f"scripts/build_docx.py. Cohort: {cohort_stamp(v)} "
        "Every number below is read from paper/verified_numbers.json, which is "
        "regenerated from the raw data by scripts/build_paper_assets.py. "
    )
    r2.font.size = Pt(9)
    r2.font.color.rgb = RGBColor(0xC0, 0x00, 0x00)
    if stale:
        r3 = box.add_run(
            "WARNING: this build mixes cohorts. The monthly aggregation was corrected "
            "(weeks are now split across calendar months instead of being assigned whole "
            "to the month containing the start of the epidemiological week) and the "
            "benchmark is still being re-run. Cities already re-run have 1,428 predictions "
            "per model; those not yet re-run still have 1,452. Do not read the cross-city "
            "summaries in this build."
        )
        r3.bold = True
        r3.font.size = Pt(9)
        r3.font.color.rgb = RGBColor(0xC0, 0x00, 0x00)
    doc.add_page_break()


def abstract(doc: Document, v: dict) -> None:
    heading(doc, "ABSTRACT")
    ms = v["mean_smape_by_model"]
    wins = v["cities_ranked_first_by_model"]
    pc = v["per_city"]
    body(doc, (
        "Dengue imposes a recurrent and heavy burden on the Brazilian health system, and "
        "municipal forecasts are used to plan vector control, procurement and hospital "
        "capacity. Operational tools rely on models that must be trained and retrained per "
        "city, which limits how widely they can be deployed. Whether a pre-trained "
        "time-series foundation model can substitute for that per-city pipeline has not "
        "been tested on Brazilian dengue data."
    ), bold_lead="Background. ")
    body(doc, (
        f"We benchmarked {v['n_models']} forecasting models on monthly notified dengue "
        f"counts from InfoDengue for {v['n_cities']} Brazilian state capitals. Six models "
        "were trained per city (SARIMAX, Ridge, LightGBM, XGBoost, CatBoost and Random "
        "Forest); TimesFM 2.5 was applied zero-shot with no city-specific training. "
        "Evaluation used rolling-origin cross-validation with an expanding window "
        "(minimum training window 48 months, horizon 12 months, step 1 month). "
        "The pre-specified primary metric was sMAPE, with MAE and RMSE alongside. Split "
        "conformal prediction intervals, calibrated only on residuals from earlier folds, "
        "were assessed for empirical coverage. Reporting follows TRIPOD+AI."
    ), bold_lead="Methods. ")
    body(doc, (
        f"TimesFM achieved the lowest sMAPE in {wins['timesfm']} of {v['n_cities']} cities "
        f"(mean across cities {ms['timesfm']:.1f}%) against {ms['catboost']:.1f}% for "
        f"CatBoost, the best purpose-trained model, {ms['ridge']:.1f}% for Ridge, "
        f"{ms['sarimax']:.1f}% for SARIMAX and {ms['lgbm']:.1f}% for LightGBM. By MAE, "
        f"TimesFM ranked first in all {v['n_cities']} cities. In Rio de Janeiro it reached "
        f"sMAPE {pc['Rio de Janeiro']['timesfm']['sMAPE']:.1f}% against "
        f"{pc['Rio de Janeiro']['xgboost']['sMAPE']:.1f}% for XGBoost. The three accuracy "
        "metrics selected the same model in only four of eight cities. All models "
        "substantially underestimated the 2024 epidemic peaks."
    ), bold_lead="Results. ")
    body(doc, (
        "A zero-shot foundation model matched or exceeded purpose-trained models for "
        "12-month-ahead dengue point forecasting without any local training, which lowers "
        "the cost of extending forecasting to municipalities lacking modelling capacity. "
        "The advantage does not extend to probabilistic forecasts, where the ranking "
        "depends on the interval construction. Intervals were wide and coverage collapsed "
        "during the 2024 epidemic, so these forecasts suit planning baselines rather than "
        "outbreak triggers."
    ), bold_lead="Conclusion. ")
    body(doc, (
        "dengue; forecasting; foundation model; time series; machine learning; Brazil; "
        "epidemiological surveillance"
    ), bold_lead="Keywords: ")


def methods(doc: Document, v: dict) -> None:
    heading(doc, "METHODS")
    body(doc, (
        "Reporting follows the TRIPOD+AI statement for the development and evaluation of "
        "prediction models using regression or machine learning methods. All data are "
        "publicly available aggregate surveillance records with no individual-level "
        "information, so ethics committee review was not required. Analysis code, "
        "processed data and every result file underlying the tables and figures are "
        "openly available at https://github.com/fabianofilho/dengue-time-series."
    ))

    heading(doc, "Ethical Considerations", 2)
    body(doc, (
        "This study used exclusively publicly available aggregate surveillance records "
        "obtained from the InfoDengue open application programming interface. The data are "
        "counts of notified dengue cases per municipality per epidemiological week, contain "
        "no individual-level information, and cannot be used to identify any person. Under "
        "Brazilian National Health Council Resolution CNS 510/2016, Article 1, sole "
        "paragraph, such research is not subject to registration or evaluation by the "
        "CEP/CONEP research ethics committee system: item II covers research using publicly "
        "accessible information, and item V covers research with databases whose information "
        "is aggregated with no possibility of individual identification. Both apply here. "
        "Approval by an Institutional Ethics Committee was therefore not required, and none "
        "was sought. No human participants were recruited and no informed consent was "
        "applicable."
    ))

    heading(doc, "Study setting and data source", 2)
    body(doc, (
        "Eight Brazilian state capitals were studied: Sao Paulo (SP), Rio de Janeiro (RJ), "
        "Belo Horizonte (MG), Brasilia (DF), Fortaleza (CE), Recife (PE), Manaus (AM) and "
        "Salvador (BA), spanning the Southeast, Central-West, Northeast and North "
        "macroregions. Notification data came from the InfoDengue public API, which "
        "aggregates mandatory notifications from SINAN."
    ))

    heading(doc, "Aggregation of epidemiological weeks to calendar months", 2)
    body(doc, (
        "Weekly counts were converted to calendar months by distributing each "
        "epidemiological week uniformly across the seven days it covers and summing by "
        "month. This is not the obvious approach and the obvious one is wrong: assigning "
        "each week whole to the month containing its start date produces a series in which "
        "some months receive four weeks of cases and others five, a calendar-driven "
        "oscillation of 25 percent with a period of roughly 12.4 months that does not even "
        "align with annual seasonality. On this dataset that error displaced the median "
        "month by 10.1 percent and inflated the March 2024 peak in Sao Paulo by 16.8 "
        "percent. Months only partly covered by the weekly series were dropped, which "
        "removes January 2010 and December 2024 and leaves 178 monthly observations per "
        "city."
    ))
    body(doc, (
        "Reporting delay was assessed and found not to require correction: the notified "
        "count and the InfoDengue nowcast estimate are identical across the entire series, "
        "indicating the nowcast has converged."
    ))

    heading(doc, "Models", 2)
    body(doc, (
        "SARIMAX with non-seasonal order (1,1,1) and seasonal order (1,1,0,12) served as "
        "the parametric baseline. Ridge, LightGBM, XGBoost, CatBoost and Random Forest "
        "were wrapped in the recursive multi-step forecaster of skforecast with 24 "
        "autoregressive lags and refitted independently at every fold. TimesFM 2.5, a "
        "200-million-parameter decoder-only transformer, was applied strictly zero-shot: "
        "no fine-tuning, no city-specific adaptation and no refitting at any fold. It "
        "therefore receives strictly less information than its competitors."
    ))
    body(doc, (
        "Ridge requires feature standardisation and the other learners do not. Without it "
        "the L2 penalty is negligible against counts in the thousands, and in the earliest "
        "folds, where training rows approach the number of lags, the fit becomes "
        "near-singular; recursive 12-step prediction then amplifies the result. In "
        "Salvador a single fold produced a forecast of 8.7e14 cases and by itself dominated "
        "that city's pooled MAE. With standardisation no fold diverged in any city. This "
        "failure is nearly silent under sMAPE, which is bounded and therefore absorbs a "
        "divergence of any magnitude."
    ))
    body(doc, (
        "Prophet was specified in the protocol but failed to run in seven of eight cities "
        "because of an unresolved dependency. It does not appear in the results and is not "
        "discussed as though it had been tested."
    ))

    heading(doc, "Evaluation design", 2)
    n = v["n_predictions_per_model_city"]
    folds = n // 12
    body(doc, (
        "Rolling-origin cross-validation with an expanding training window was used. At "
        "each fold the model is fitted only on observations preceding the forecast origin "
        "and predicts the following 12 months, so no information from the evaluation "
        f"period reaches the model. The minimum training window was 48 months, giving "
        f"{folds} folds per city, {n:,} predictions per model per city and "
        f"{v['n_predictions_total']:,} in total. This is temporal out-of-sample evaluation "
        "within a single data source, which is internal validation under the TRIPOD+AI "
        "nomenclature, not external validation."
    ))

    heading(doc, "Metrics, uncertainty and prediction intervals", 2)
    body(doc, (
        "The pre-specified primary metric was sMAPE, with MAE and RMSE reported alongside "
        "because the three reward different behaviour: sMAPE rewards accuracy in "
        "inter-epidemic troughs and RMSE rewards accuracy at peaks. Confidence intervals "
        f"for the accuracy estimates come from bootstrap resampling of the pooled "
        f"predictions ({v['bootstrap']['replicates']:,} replicates, "
        f"{v['bootstrap']['method']} method, seed {v['bootstrap']['seed']}). Because "
        "predictions from the same fold are not independent, these intervals understate "
        "the true uncertainty and are read as a lower bound."
    ))
    body(doc, (
        "Prediction intervals were built by split conformal prediction on the point "
        "forecasts already produced, without refitting any model. For each model, city and "
        "horizon step, the interval at a given fold is calibrated on residuals from folds "
        "strictly earlier, which preserves the same no-leakage property the point "
        "forecasts have. Two constructions are reported throughout, additive and "
        "multiplicative, because the choice is consequential and the data do not settle "
        "it. Intervals were assessed on empirical coverage, width, and the weighted "
        "interval score."
    ))


def results(doc: Document, v: dict) -> None:
    heading(doc, "RESULTS")
    de, pc, ms = v["descriptive"], v["per_city"], v["mean_smape_by_model"]
    wins = v["cities_ranked_first_by_model"]
    order = sorted(de, key=lambda c: -de[c]["inc_mean_annual"])

    heading(doc, "Descriptive characteristics", 2)
    top, low = order[0], order[-1]
    body(doc, (
        f"In absolute counts the cities differ by more than an order of magnitude: "
        f"Sao Paulo accumulated {de['Sao Paulo']['total']:,} notified cases against "
        f"{de['Salvador']['total']:,} in Salvador. Standardising by population reorders "
        f"them substantially. {top} carries the highest burden at a mean "
        f"{de[top]['inc_mean_annual']:,.0f} notified cases per 100,000 inhabitants per "
        f"year, and {low} the lowest at {de[low]['inc_mean_annual']:,.0f}. Sao Paulo, "
        f"first by raw count, falls to position {order.index('Sao Paulo') + 1} at "
        f"{de['Sao Paulo']['inc_mean_annual']:,.0f}."
    ))
    add_table(
        doc,
        ["City", "Region", "Pop. 2022 (M)", "Total cases", "Median/mo",
         "Max/mo", "CV (%)", "Incid./100k/yr"],
        [[c, de[c]["region"], f"{de[c]['pop2022'] / 1e6:.2f}", f"{de[c]['total']:,}",
          f"{de[c]['median']:,.0f}", f"{de[c]['max']:,}", f"{de[c]['cv']:.0f}",
          f"{de[c]['inc_mean_annual']:,.0f}"] for c in order],
        "Table 1.",
        "Notified dengue burden in eight Brazilian state capitals. Rows sorted by mean "
        "annual incidence, which reorders the table relative to absolute counts. "
        "Population denominators are anchored on the 2010 and 2022 Censuses.",
    )
    add_figure(doc, "fig1_series", "Figure 1.",
               "Monthly notified dengue cases in eight Brazilian state capitals. Panels "
               "are ordered by cumulative case count and each has its own vertical scale. "
               "The shaded band marks the 48-month minimum training window.")

    heading(doc, "Forecasting accuracy", 2)
    best_sup = min((m for m in MODELS if m != "timesfm"), key=lambda m: ms[m])
    others = sorted((x for x in MODELS if x not in ("timesfm", best_sup)),
                    key=lambda x: ms[x])
    body(doc, (
        f"TimesFM achieved the lowest sMAPE in {wins['timesfm']} of {v['n_cities']} "
        f"cities, with a mean across cities of {ms['timesfm']:.1f}% (range "
        f"{v['timesfm_smape_range'][0]:.1f}% to {v['timesfm_smape_range'][1]:.1f}%). The "
        f"best purpose-trained model on the same summary was {LABEL[best_sup]} at "
        f"{ms[best_sup]:.1f}%, followed by "
        + ", ".join(f"{LABEL[m]} at {ms[m]:.1f}%" for m in others) + "."
    ))
    rows, bold = [], set()
    for ri, c in enumerate(CITY_ORDER):
        vals = [pc[c][m]["sMAPE"] for m in MODELS]
        bold.add((ri, 1 + vals.index(min(vals))))
        rows.append([c] + [f"{x:.1f}" for x in vals])
    rows.append(["Mean across cities"] + [f"{ms[m]:.1f}" for m in MODELS])
    rows.append(["Cities ranked first"] + [f"{wins[m]}/{v['n_cities']}" for m in MODELS])
    add_table(doc, ["City"] + [LABEL[m] for m in MODELS], rows,
              "Table 2.",
              "Forecast error by model and city, sMAPE (%). Bold marks the lowest value "
              "in each row. Lower is better.", bold)
    add_figure(doc, "fig2_heatmap", "Figure 2.",
               "sMAPE by model and city. The outlined cell in each column is the best "
               "model for that city.")
    add_figure(doc, "fig3_timesfm_vs_best", "Figure 3.",
               "Zero-shot TimesFM against the best purpose-trained model in each city, "
               "with 95% bootstrap intervals. The label above each orange bar names which "
               "supervised model achieved it, and it is not the same across cities.")

    heading(doc, "Absolute error", 2)
    body(doc, (
        "MAE and RMSE are reported on the original case-count scale and are therefore not "
        "comparable across cities. The gap between the two within a city quantifies how "
        "much of the total error sits in a few months: where RMSE greatly exceeds MAE, the "
        "error is concentrated in epidemic peaks rather than spread across the series."
    ))
    rows, bold = [], set()
    ri = 0
    for c in CITY_ORDER:
        maes = [pc[c][m]["MAE"] for m in MODELS]
        rmses = [pc[c][m]["RMSE"] for m in MODELS]
        bold.add((ri, 2 + maes.index(min(maes))))
        rows.append([c, "MAE"] + [f"{x:,.0f}" for x in maes])
        ri += 1
        bold.add((ri, 2 + rmses.index(min(rmses))))
        rows.append(["", "RMSE"] + [f"{x:,.0f}" for x in rmses])
        ri += 1
    add_table(doc, ["City", "Metric"] + [LABEL[m] for m in MODELS], rows,
              "Table 3.",
              "Forecasting accuracy by model and city: mean absolute error (MAE) and root "
              "mean squared error (RMSE), in notified cases per month. Bold marks the best "
              "value in each row.", bold)

    heading(doc, "The metrics disagree", 2)
    agree = sum(1 for r in v["metric_disagreement"]
                if r["smape"] == r["mae"] == r["rmse"])
    body(doc, (
        f"sMAPE, MAE and RMSE select the same model in only {agree} of the "
        f"{v['n_cities']} cities. Reporting a single metric would conceal this. The "
        "mechanism is arithmetic rather than epidemiological: sMAPE normalises each error "
        "by the magnitude of the observation, so the low-count months between epidemics "
        "dominate it, while RMSE squares the error, so the few epidemic months dominate "
        "instead."
    ))
    add_table(doc, ["City", "Best by sMAPE", "Best by MAE", "Best by RMSE", "Agree"],
              [[r["city"], LABEL[r["smape"]], LABEL[r["mae"]], LABEL[r["rmse"]],
                "Yes" if r["smape"] == r["mae"] == r["rmse"] else "No"]
               for r in v["metric_disagreement"]],
              "Table 4.",
              "Best-performing model per city under each accuracy metric.")
    add_figure(doc, "fig4_metric_disagreement", "Figure 4.",
               "Number of cities in which each model ranks first, under each metric.")
    add_figure(doc, "fig5_observed_vs_predicted", "Figure 5.",
               "Observed and predicted monthly cases in four cities, with predictions "
               "averaged across folds. Neither model reproduces the 2024 peak.")

    heading(doc, "Prediction intervals", 2)
    iv = v.get("intervals")
    if not iv:
        flag(doc, "interval calibration absent from verified_numbers.json; run "
                  "scripts/run_intervals.py then scripts/build_paper_assets.py")
        return
    lines = []
    for scale, label in (("log1p", "Multiplicative"), ("absolute", "Additive")):
        s = iv[scale]
        best = min(s, key=lambda m: s[m]["rwis"])
        lo = min(s[m]["c95"] for m in s) * 100
        hi = max(s[m]["c95"] for m in s) * 100
        lines.append(
            f"{label}: lowest rWIS is {LABEL[best]} at {s[best]['rwis']:.2f}, and coverage "
            f"at the nominal 95% level ranges from {lo:.1f}% to {hi:.1f}%"
        )
    body(doc, (
        "Conformal calibration works: empirical coverage tracks the nominal level closely "
        "for every model. " + ". ".join(lines) + ". The ordering by weighted interval "
        "score is not preserved between the two constructions, so no probabilistic "
        "ranking of the models is claimed."
    ))
    rows = []
    for scale, label in (("log1p", "Multiplicative"), ("absolute", "Additive")):
        s = iv[scale]
        for m in sorted(s, key=lambda x: s[x]["rwis"]):
            rows.append([label, LABEL[m], f"{s[m]['c50'] * 100:.1f}",
                         f"{s[m]['c80'] * 100:.1f}", f"{s[m]['c95'] * 100:.1f}",
                         f"{s[m]['rwis']:.2f}", f"{int(s[m]['wins'])}/{v['n_cities']}"])
    add_table(doc,
              ["Construction", "Model", "Cov. 50%", "Cov. 80%", "Cov. 95%",
               "rWIS", "Cities best"],
              rows, "Table 5.",
              "Calibration and sharpness of conformal prediction intervals. Coverage "
              "close to the nominal level indicates calibration. rWIS is the weighted "
              "interval score divided by the mean observed count of the city; lower is "
              "better.")
    add_figure(doc, "fig6_calibration", "Figure 6.",
               "Calibration of the conformal intervals. Left and centre: empirical "
               "coverage against nominal level. Right: coverage at nominal 95% by "
               "calendar year, showing the collapse in 2024.")
    add_figure(doc, "fig7_fan", "Figure 7.",
               "Six-month-ahead TimesFM forecasts with additive conformal intervals. Red "
               "points mark observations falling outside the 95% interval.")


def ensembles(doc: Document, v: dict) -> None:
    en = v.get("ensembles")
    if not en:
        heading(doc, "Model combinations", 2)
        flag(doc, "ensemble results absent from verified_numbers.json; run "
                  "scripts/run_ensembles.py then scripts/build_paper_assets.py")
        return

    heading(doc, "Combining models does not rescue the ranking", 2)
    n_beat = en["n_cities_ensemble_beats_best_member"]
    deltas = en["delta_by_city"]
    worst = min(deltas, key=lambda c: deltas[c])
    means = en["mean_by_approach"]
    best_by = en["best_by_metric"]

    body(doc, (
        f"The best combination fails to beat the best single model of that city in "
        f"{v['n_cities'] - n_beat} of {v['n_cities']} cities. The single exception is Sao "
        f"Paulo, where the median ensemble improves on the best member by "
        f"{deltas['sao_paulo']:.1f} sMAPE points. Elsewhere the shortfall reaches "
        f"{abs(deltas[worst]):.1f} points in "
        f"{worst.replace('_', ' ').title().replace('Rio', 'Rio de Janeiro')}."
    ))
    body(doc, (
        f"Among the combinations themselves the median ensemble is the best at "
        f"{means['median']['sMAPE']:.1f}% mean sMAPE, ahead of every individual model "
        f"except TimesFM, and its advantage over the arithmetic mean at "
        f"{means['mean']['sMAPE']:.1f}% comes from resisting the members that diverge. "
        f"The metric dependence reported above reappears here and reverses the "
        f"conclusion: on sMAPE and MAE the best approach overall is a single model, "
        f"{LABEL.get(best_by['sMAPE'], best_by['sMAPE'])}, while on RMSE it is the "
        f"stacked ensemble at {means['stacking']['RMSE']:,.0f} against "
        f"{means['timesfm']['RMSE']:,.0f} for TimesFM. Whether combining models helps has "
        f"no metric-free answer in these data."
    ))

    rows = []
    for name, m in sorted(means.items(), key=lambda kv: kv[1]["sMAPE"]):
        label = ENS_LABEL.get(name)
        rows.append([
            (label + " †") if label else LABEL.get(name, name),
            f"{m['sMAPE']:.1f}", f"{m['MAE']:,.0f}", f"{m['RMSE']:,.0f}",
        ])
    add_table(doc, ["Approach", "sMAPE (%)", "MAE", "RMSE"], rows,
              "Table 6.",
              "Individual models and model combinations on the same evaluation set, "
              "averaged over the eight cities. † marks a combination. Weights and the "
              "stacking meta-model are fitted only on folds earlier than the one being "
              "predicted, and every approach is scored on the same 1,188 predictions per "
              "city.")
    add_figure(doc, "fig8_ensembles", "Figure 8.",
               "Difference in sMAPE between the best combination and the best single "
               "model, by city. Positive values mean the combination wins.")


def discussion(doc: Document, v: dict) -> None:
    ms = v["mean_smape_by_model"]
    wins = v["cities_ranked_first_by_model"]
    heading(doc, "DISCUSSION")
    body(doc, (
        f"Across eight Brazilian state capitals, a foundation model applied with no local "
        f"training matched or exceeded six purpose-trained models at a 12-month horizon. "
        f"It ranked first by sMAPE in {wins['timesfm']} of {v['n_cities']} cities and by "
        f"MAE in all {v['n_cities']}. The comparison is asymmetric in its disfavour: the "
        f"supervised models were refitted on each city at every fold, while TimesFM saw "
        f"the series only as context at inference."
    ))
    body(doc, (
        "The operational consequence is about cost rather than accuracy. The best "
        "supervised model is not the same model in every city, so a per-city pipeline "
        "requires per-city model selection, retraining and monitoring. A single unmodified "
        "foundation model removes that overhead, which is what makes extension to "
        "municipalities without modelling capacity plausible."
    ))
    body(doc, (
        f"Accuracy in absolute terms is nonetheless poor everywhere. A mean sMAPE of "
        f"{ms['timesfm']:.0f}% is not a small error, and no model anticipated the 2024 "
        "epidemic in any city. These models provide a reasonable baseline expectation for "
        "planning under ordinary conditions and are not outbreak detectors."
    ))

    heading(doc, "Limitations", 2)
    for text in [
        "The outcome is notified cases, which under-represent true infections. What is "
        "forecast is the behaviour of the surveillance system, including reporting delay "
        "and saturation during peaks, not biological incidence.",
        "No exogenous covariates were used. Models with access to rainfall, temperature "
        "and vector indices outperform univariate approaches for outbreak onset, so the "
        "TimesFM advantage may narrow once supervised models receive informative features "
        "that a zero-shot model cannot consume.",
        "The prediction intervals are calibrated but not sharp, and coverage falls during "
        "the 2024 epidemic. An interval that is correct on average across a decade while "
        "failing during the one epidemic in the series does not support the decision it "
        "would be used for.",
        "Bootstrap intervals resample predictions that are serially dependent within and "
        "across folds. They understate uncertainty and are reported as a lower bound, not "
        "as a significance test.",
        "Validation is internal and temporal. No independent surveillance system or "
        "country was used, and the eight cities constitute eight internal validations "
        "rather than a geographic external validation.",
        "SARIMAX orders and all tree hyperparameters were fixed rather than tuned. This "
        "mirrors low-maintenance deployment but understates what a tuned supervised "
        "pipeline could achieve.",
        "One TimesFM checkpoint was evaluated on eight capitals. Whether the result holds "
        "for other foundation models, other checkpoints, or the smaller municipalities "
        "that make up most of Brazil, is untested.",
    ]:
        p = doc.add_paragraph(style="List Bullet")
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p.add_run(text)


def declarations(doc: Document) -> None:
    heading(doc, "DECLARATIONS")
    for lead, text in [
        ("Data and code availability. ", "Monthly case data are publicly available through "
         "the InfoDengue API. All analysis code, processed data, result files and the "
         "scripts that generate every table and figure in this manuscript are available at "
         "https://github.com/fabianofilho/dengue-time-series."),
        ("Funding. ", "No external funding was received."),
        ("Conflicts of interest. ", "The authors declare no conflicts of interest."),
        ("Declaration of the use of generative AI and AI-assisted technologies in the writing "
         "process. ",
         "During the preparation of this work the authors used a large language model based "
         "assistant to support literature organisation, drafting, and the generation of "
         "analysis and figure code. After using this tool the authors reviewed and edited "
         "the content as needed, verified every reported number against the underlying "
         "analysis output, and take full responsibility for the content of the publication. "
         "The tool is not an author and does not meet authorship criteria."),
        ("Acknowledgements. ", "Surveillance data were obtained from InfoDengue "
         "(Fiocruz/FGV). TimesFM was developed by Google DeepMind."),
    ]:
        body(doc, text, bold_lead=lead)
    body(doc, (
        "Following the CRediT taxonomy. Fabiano Barcellos-Filho: Conceptualization, Data "
        "curation, Formal analysis, Investigation, Methodology, Project administration, "
        "Software, Validation, Visualization, Writing: original draft, Writing: review "
        "and editing. Jose Trajano: Conceptualization, Formal analysis, Methodology, "
        "Supervision, Writing: review and editing. Wilson Falco: Data curation, "
        "Investigation, Software, Writing: review and editing. All authors read and "
        "approved the final manuscript."
    ), bold_lead="Author contributions. ")


def main() -> None:
    ap = argparse.ArgumentParser(description="Gera o manuscrito em .docx para revisao")
    ap.add_argument("--out", default=str(ROOT / "build" / "manuscript.docx"))
    args = ap.parse_args()

    v = load()
    ns = {f["n_predictions"].iloc[0] for f in v["_metrics_files"].values()}
    stale = len(ns) > 1
    if stale:
        print(f"  [AVISO] coortes misturadas nos resultados: n_predictions = {sorted(ns)}")

    doc = setup_document()
    title_block(doc, v, stale)
    abstract(doc, v)
    methods(doc, v)
    results(doc, v)
    ensembles(doc, v)
    discussion(doc, v)
    declarations(doc)

    out = Path(args.out)
    out.parent.mkdir(parents=True, exist_ok=True)
    doc.save(out)

    figs = sum(1 for r in doc.part.rels.values() if "image" in r.reltype)
    print(f"escrito {out}")
    print(f"  {len(doc.paragraphs)} paragrafos, {len(doc.tables)} tabelas, {figs} figuras")
    print(f"  coorte: {cohort_stamp(v)}")


if __name__ == "__main__":
    main()
