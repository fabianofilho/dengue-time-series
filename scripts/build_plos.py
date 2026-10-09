"""
Monta o manuscrito de submissao (PLOS NTD) e o .docx de revisao a partir de UMA fonte.

Por que existe: ate 06/10/2026 o texto vivia em dois lugares, paper/manuscript.tex
(parado em 23/09, com as citacoes) e scripts/build_docx.py (o que os coautores
revisaram, sem citacoes no Methods nem na Discussion). Os dois divergiram. Agora o
texto mora so em paper/plos/manuscript_src.tex e os dois artefatos saem dele:

  manuscript_src.tex ──> build/plos/manuscript.tex  (arquivo unico, bibliografia
                     │                               embutida, como a PLOS pede)
                     ├──> build/plos/manuscript.pdf  (submissao: sem figuras, so
                     │                               legendas, como a PLOS pede)
                     ├──> build/plos/manuscript_review.docx (revisao: figuras no
                     │                               ponto do texto, via pandoc)
                     ├──> build/plos/figures/FigN.tif (TIFF 300 dpi, RGB)
                     └──> build/plos/S1_Appendix.pdf

Nenhum numero e digitado no .tex. Os marcadores:

  \\V{expr}            expressao Python avaliada contra verified_numbers.json (nome v),
                      com a sintaxe de campo de f-string: \\V{v['n_cities']},
                      \\V{v['mean_smape_by_model']['timesfm']:.1f}. Erro interrompe.
  \\plosfig{arq}{rotulo}{titulo}{legenda}
                      figura: na submissao so a legenda; na revisao, imagem + legenda
  \\pending{texto}     pendencia em vermelho; o build avisa quantas restam
  \\reviewonly{texto}  so aparece no .docx de revisao

Citacao e referencia cruzada no .docx saem resolvidas pelo .aux da compilacao do
PDF, entao a numeracao das referencias e a mesma nos dois artefatos.

Uso:
    PYTHONPATH=src python scripts/build_paper_assets.py   # antes, se os dados mudaram
    python scripts/build_plos.py
"""
from __future__ import annotations

import json
import os
import re
import shutil
import subprocess
import sys
from datetime import datetime, timezone
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
PAPER = ROOT / "paper"
PLOS = PAPER / "plos"
OUT = ROOT / "build" / "plos"
# TECTONIC no ambiente, senao o do PATH, senao o env conda do Mac do Fabiano
TECTONIC = (os.environ.get("TECTONIC") or shutil.which("tectonic")
            or "/opt/miniconda3/envs/tectonic/bin/tectonic")

LABEL = {
    "timesfm": "TimesFM", "catboost": "CatBoost", "xgboost": "XGBoost",
    "randomforest": "Random Forest", "sarimax": "SARIMAX", "ridge": "Ridge",
    "lgbm": "LightGBM", "seasonal_naive": "seasonal naive",
    "global_randomforest": "global Random Forest", "global_catboost": "global CatBoost",
    "global_lgbm": "global LightGBM", "global_xgboost": "global XGBoost",
    "global_ridge": "global Ridge",
    "mean": "mean ensemble", "median": "median ensemble",
    "weighted": "weighted ensemble", "stacking": "stacked ensemble",
}
TRAINED = ["catboost", "xgboost", "randomforest", "sarimax", "ridge", "lgbm"]
WORD = {0: "zero", 1: "one", 2: "two", 3: "three", 4: "four", 5: "five", 6: "six",
        7: "seven", 8: "eight", 9: "nine", 10: "ten", 11: "eleven", 12: "twelve"}
MONTH = {"01": "January", "02": "February", "03": "March", "04": "April", "05": "May",
         "06": "June", "07": "July", "08": "August", "09": "September", "10": "October",
         "11": "November", "12": "December"}
CITY_TEX = {"Sao Paulo": "S\\~ao Paulo", "Brasilia": "Bras\\'ilia"}


# --------------------------------------------------------------------------- #
# funcoes disponiveis dentro de \V{...}
# --------------------------------------------------------------------------- #
def periodo(v: dict) -> str:
    ini, fim = v["study_period"].split(" to ")
    (ai, mi), (af, mf) = ini.split("-"), fim.split("-")
    return f"{MONTH[mi]} {ai} to {MONTH[mf]} {af}"


def city(nome: str) -> str:
    return CITY_TEX.get(nome, nome)


def lista(itens: list[str]) -> str:
    itens = list(itens)
    if len(itens) <= 1:
        return "".join(itens)
    return ", ".join(itens[:-1]) + " and " + itens[-1]


def lead(v: dict, cidade: str, metrica: str) -> str:
    """Modelo com o menor valor da metrica na cidade, entre os sete publicados."""
    pc = v["per_city"][cidade]
    return min(pc, key=lambda m: pc[m][metrica])


def agree(v: dict) -> int:
    return sum(1 for r in v["metric_disagreement"] if r["smape"] == r["mae"] == r["rmse"])


def md(v: dict, cidade: str) -> dict:
    """Melhor modelo por metrica na cidade (linha de metric_disagreement)."""
    return next(r for r in v["metric_disagreement"] if r["city"] == cidade)


def ens_worst(v: dict) -> tuple[str, float]:
    """Cidade em que a melhor combinacao fica mais atras do melhor modelo isolado."""
    d = v["ensembles"]["delta_by_city"]
    slug = min(d, key=lambda c: d[c])
    # os slugs dos ensembles nao seguem um padrao so ("rio", "sao_paulo")
    nome = next(c for c in v["per_city"] if c.lower().replace(" ", "_").startswith(slug))
    return nome, abs(d[slug])


def wins_by(v: dict, metrica: str, modelo: str = "timesfm") -> int:
    """Cidades em que o modelo e o melhor pela metrica (smape, mae ou rmse)."""
    return sum(1 for r in v["metric_disagreement"] if r[metrica] == modelo)


def namespace(v: dict) -> dict:
    return {"v": v, "L": LABEL, "W": WORD, "TRAINED": TRAINED, "periodo": periodo,
            "city": city, "lista": lista, "lead": lead, "agree": agree, "md": md,
            "ens_worst": ens_worst, "wins_by": wins_by,
            "min": min, "max": max, "len": len, "sum": sum, "abs": abs, "round": round,
            "sorted": sorted, "set": set, "format": format, "next": next}


# --------------------------------------------------------------------------- #
# parser de argumentos com chaves balanceadas
# --------------------------------------------------------------------------- #
def _arg(s: str, i: int) -> tuple[str, int]:
    """Le um argumento {...} a partir de s[i] == '{'; devolve (conteudo, fim)."""
    if s[i] != "{":
        raise ValueError(f"esperava '{{' na posicao {i}: {s[i:i + 40]!r}")
    depth, j = 0, i
    while j < len(s):
        c = s[j]
        if c == "\\":
            j += 2
            continue
        if c == "{":
            depth += 1
        elif c == "}":
            depth -= 1
            if depth == 0:
                return s[i + 1:j], j + 1
        j += 1
    raise ValueError(f"chave nao fechada a partir de {s[i:i + 60]!r}")


def macro(s: str, nome: str, nargs: int, fn) -> str:
    out, i = [], 0
    tag = "\\" + nome + "{"
    while True:
        k = s.find(tag, i)
        if k < 0:
            out.append(s[i:])
            return "".join(out)
        out.append(s[i:k])
        j = k + len(tag) - 1
        args = []
        for _ in range(nargs):
            a, j = _arg(s, j)
            args.append(a)
        out.append(fn(*args))
        i = j


# --------------------------------------------------------------------------- #
# render
# --------------------------------------------------------------------------- #
def render(src: str, v: dict, mode: str) -> tuple[str, list[str], list[str]]:
    """mode: 'submit' ou 'review'. Devolve (tex, pendencias, figuras na ordem)."""
    ns = namespace(v)

    def val(expr: str) -> str:
        try:
            # ns vai nos globais: compreensoes e lambdas nao enxergam os locais do eval
            return eval('f"{' + expr + '}"', {"__builtins__": {}, **ns})
        except Exception as e:  # noqa: BLE001
            raise SystemExit(f"\\V{{{expr}}} falhou: {e}") from e

    tex = macro(src, "V", 1, val)

    pend: list[str] = []

    def pending(t: str) -> str:
        pend.append(re.sub(r"\s+", " ", t).strip())
        return "\\textcolor{red}{\\textbf{[PENDING: " + t + "]}}"

    tex = macro(tex, "pending", 1, pending)
    tex = macro(tex, "reviewonly", 1, lambda t: t if mode == "review" else "")

    figs: list[str] = []

    def fig(arq: str, rotulo: str, titulo: str, legenda: str) -> str:
        figs.append(arq)
        img = ""
        if mode == "review":
            img = f"\\includegraphics[width=\\textwidth]{{{(PAPER / 'figures' / arq).as_posix()}.png}}\n"
        return (f"\\begin{{figure}}[!h]\n{img}\\caption{{\\textbf{{{titulo}}} {legenda}}}\n"
                f"\\label{{{rotulo}}}\n\\end{{figure}}")

    tex = macro(tex, "plosfig", 4, fig)
    tex = re.sub(r"\\input\{tables/([^}]+)\}",
                 lambda m: (PAPER / "tables" / f"{m.group(1)}.tex").read_text(encoding="utf-8"),
                 tex)
    return tex, pend, figs


# --------------------------------------------------------------------------- #
# compilacao
# --------------------------------------------------------------------------- #
def tectonic(tex: Path) -> None:
    r = subprocess.run([TECTONIC, "-X", "compile", tex.name, "--keep-intermediates"],
                       cwd=tex.parent, capture_output=True, text=True)
    (tex.parent / f"{tex.stem}.compile.log").write_text(r.stdout + r.stderr)
    if r.returncode != 0:
        print((r.stdout + r.stderr)[-3000:])
        raise SystemExit(f"tectonic falhou em {tex.name}")


def aux_maps(aux: Path) -> tuple[dict[str, int], dict[str, str]]:
    txt = aux.read_text(encoding="utf-8")
    cites = {k: int(n) for k, n in re.findall(r"\\bibcite\{([^}]+)\}\{(\d+)\}", txt)}
    labels = {k: n for k, n in re.findall(r"\\newlabel\{([^}]+)\}\{\{([^}]*)\}", txt)}
    return cites, labels


def resolve_for_docx(tex: str, cites: dict[str, int], labels: dict[str, str]) -> str:
    def cite(m: re.Match) -> str:
        ns_ = sorted(cites[k.strip()] for k in m.group(1).split(","))
        return "[" + ", ".join(str(n) for n in ns_) + "]"

    # o pandoc nao numera legenda no Word; o numero vem do .aux, igual ao PDF
    out, i = [], 0
    while (k := tex.find("\\caption{", i)) >= 0:
        corpo, fim = _arg(tex, k + len("\\caption"))
        lab = re.match(r"\s*\\label\{([^}]+)\}", tex[fim:])
        prefixo = ""
        if lab and lab.group(1) in labels:
            tipo = "Table" if lab.group(1).startswith("tab:") else "Fig"
            prefixo = f"\\textbf{{{tipo} {labels[lab.group(1)]}.}} "
        out.append(tex[i:k] + "\\caption{" + prefixo + corpo + "}")
        i = fim
    tex = "".join(out) + tex[i:]
    tex = re.sub(r"\\cite\{([^}]+)\}", cite, tex)
    tex = re.sub(r"\\ref\{([^}]+)\}", lambda m: labels[m.group(1)], tex)
    return tex


def export_tiffs(figs: list[str]) -> list[Path]:
    from PIL import Image

    dest = OUT / "figures"
    dest.mkdir(parents=True, exist_ok=True)
    out = []
    for n, arq in enumerate(figs, 1):
        im = Image.open(PAPER / "figures" / f"{arq}.png")
        if im.mode in ("RGBA", "LA"):
            fundo = Image.new("RGB", im.size, "white")
            fundo.paste(im, mask=im.split()[-1])
            im = fundo
        else:
            im = im.convert("RGB")
        # PLOS: largura 789 a 2250 px e altura ate 2625 px a 300 dpi.
        escala = min(1.0, 2250 / im.width, 2625 / im.height)
        if escala < 1.0:
            im = im.resize((round(im.width * escala), round(im.height * escala)),
                           Image.LANCZOS)
        p = dest / f"Fig{n}.tif"
        im.save(p, compression="tiff_lzw", dpi=(300, 300))
        out.append(p)
    return out


def reference_docx() -> Path:
    """Estilo do .docx de revisao: Times 12, espaco duplo (pedido da PLOS)."""
    from docx import Document
    from docx.shared import Pt

    ref = OUT / "reference.docx"
    data = subprocess.run(["pandoc", "--print-default-data-file", "reference.docx"],
                          capture_output=True, check=True).stdout
    ref.write_bytes(data)
    doc = Document(ref)
    for nome in ("Normal", "Body Text", "First Paragraph", "Compact", "Bibliography"):
        if nome in [s.name for s in doc.styles]:
            st = doc.styles[nome]
            st.font.name = "Times New Roman"
            st.font.size = Pt(12)
            st.paragraph_format.line_spacing = 2.0
    for s in doc.styles:
        if s.name.startswith("Heading") or s.name in ("Title", "Caption", "Table Caption",
                                                       "Image Caption"):
            s.font.name = "Times New Roman"
    doc.save(ref)
    return ref


def postprocess_docx(path: Path) -> None:
    """Numero de linha e de pagina, e pendencias em vermelho."""
    from docx import Document
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import RGBColor

    doc = Document(path)
    sec = doc.sections[0]
    ln = OxmlElement("w:lnNumType")
    ln.set(qn("w:countBy"), "1")
    ln.set(qn("w:restart"), "continuous")
    ln.set(qn("w:distance"), "360")
    sec._sectPr.append(ln)
    p = sec.footer.paragraphs[0]
    run = p.add_run()
    for kind, text in (("begin", None), (None, "PAGE"), ("end", None)):
        if kind:
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), kind)
        else:
            el = OxmlElement("w:instrText")
            el.set(qn("xml:space"), "preserve")
            el.text = text
        run._r.append(el)

    def pinta(paragrafos):
        for par in paragrafos:
            dentro = False
            for r in par.runs:
                if "[PENDING" in r.text or "REVIEW COPY" in r.text:
                    dentro = True
                if dentro:
                    r.font.color.rgb = RGBColor(0xC0, 0x00, 0x00)
                if "]" in r.text and dentro and "REVIEW COPY" not in r.text:
                    dentro = False

    # O leitor LaTeX do pandoc compoe ' como aspa curva e -- como travessao, mesmo
    # com -smart. Regra do projeto: so aspas retas e hifen.
    troca = str.maketrans({"\u2019": "'", "\u2018": "'", "\u201c": '"', "\u201d": '"',
                           "\u2013": "-", "\u2014": "-"})

    def retas(paragrafos):
        for par in paragrafos:
            for r in par.runs:
                if r.text != r.text.translate(troca):
                    r.text = r.text.translate(troca)

    retas(doc.paragraphs)
    for t in doc.tables:
        for row in t.rows:
            for c in row.cells:
                retas(c.paragraphs)
    pinta(doc.paragraphs)
    for t in doc.tables:
        for row in t.rows:
            for c in row.cells:
                pinta(c.paragraphs)
    doc.save(path)


def main() -> None:
    v = json.loads((PAPER / "verified_numbers.json").read_text(encoding="utf-8"))
    OUT.mkdir(parents=True, exist_ok=True)
    for f in (PLOS / "plos2025.bst", PAPER / "refs" / "references.bib"):
        shutil.copy(f, OUT / f.name)

    # 1. submissao: compila com bibtex, embute o .bbl e compila de novo
    src = (PLOS / "manuscript_src.tex").read_text(encoding="utf-8")
    sub, pend, figs = render(src, v, "submit")
    stage = OUT / "manuscript.tex"
    stage.write_text(sub, encoding="utf-8")
    tectonic(stage)
    bbl = (OUT / "manuscript.bbl").read_text(encoding="utf-8")
    final = sub.replace("\\bibliography{references}", bbl)
    stage.write_text(final, encoding="utf-8")
    tectonic(stage)
    cites, labels = aux_maps(OUT / "manuscript.aux")

    # 2. S1 Appendix
    s1 = PLOS / "S1_Appendix_src.tex"
    if s1.exists():
        s1_tex, s1_pend, _ = render(s1.read_text(encoding="utf-8"), v, "submit")
        (OUT / "S1_Appendix.tex").write_text(s1_tex, encoding="utf-8")
        tectonic(OUT / "S1_Appendix.tex")
        pend += s1_pend

    # 3. figuras para upload
    tiffs = export_tiffs(figs)

    # 4. revisao: mesma fonte, figuras embutidas, citacoes resolvidas, via pandoc
    rev, _, _ = render(src, v, "review")
    # no Word a bibliografia vira lista numerada com titulo; a ordem dos \bibitem
    # e a ordem de citacao, entao o numero bate com o do PDF
    lista_refs = re.sub(r"\\begin\{thebibliography\}\{[^}]*\}",
                        lambda m: "\\section*{References}\n\\begin{enumerate}", bbl)
    lista_refs = re.sub(r"\\bibitem\{[^}]+\}", lambda m: "\\item", lista_refs)
    lista_refs = lista_refs.replace("\\end{thebibliography}", "\\end{enumerate}")
    rev = rev.replace("\\bibliography{references}", lista_refs)
    rev = resolve_for_docx(rev, cites, labels)
    stamp = datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M UTC")
    rev = rev.replace("REVIEWSTAMP", stamp)
    rev_tex = OUT / "manuscript_review.tex"
    rev_tex.write_text(rev, encoding="utf-8")
    docx = OUT / "manuscript_review.docx"
    subprocess.run(["pandoc", rev_tex.name, "-f", "latex-smart", "-t", "docx",
                    "--reference-doc", str(reference_docx()), "-o", docx.name],
                   cwd=OUT, check=True)
    postprocess_docx(docx)

    print(f"escrito {stage.relative_to(ROOT)} e manuscript.pdf")
    print(f"escrito {docx.relative_to(ROOT)}")
    print(f"  {len(cites)} referencias citadas, {len(figs)} figuras, "
          f"{len(tiffs)} TIFF em {(OUT / 'figures').relative_to(ROOT)}")
    if pend:
        print(f"  [AVISO] {len(pend)} pendencia(s) no texto:")
        for p in pend:
            print(f"    - {p[:140]}")


if __name__ == "__main__":
    sys.exit(main())
